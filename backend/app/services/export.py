"""Build a formatted .xlsx workbook from leads — mirrors the CLI script layout."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.leads import Lead

HEADER = [
    "Rank",
    "Business Name",
    "Phone",
    "Website",
    "Address",
    "Rating",
    "Reviews",
    "Quality Score",
    "Has Website",
    "Status",
    "Google Maps",
]
_WIDTHS = [6, 34, 16, 40, 46, 8, 9, 13, 12, 14, 44]


def leads_to_xlsx(service: str, location: str, leads: list[Lead]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = service[:31] or "Leads"

    header_fill = PatternFill("solid", fgColor="0E9F6E")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(HEADER)
    for col in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, lead in enumerate(leads, 1):
        ws.append(
            [
                i,
                lead.name,
                lead.phone,
                lead.website,
                lead.address,
                lead.rating if lead.rating is not None else "",
                lead.reviews,
                lead.score,
                "Yes" if lead.has_website else "No",
                lead.status,
                lead.maps_uri,
            ]
        )

    for idx, width in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{len(leads) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
